import csv
import io
import uuid
from datetime import UTC, datetime
from unittest.mock import Mock, patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.background_task import BackgroundTask
from app.models.candidate import Candidate, CandidateFormResponse
from app.models.recruiter import Recruiter
from app.models.template import TemplateField
from app.schemas.enums import TaskType
from app.services import ranking_service
from app.services.auth_service import SESSION_COOKIE, create_session_cookie
from app.tasks.batch_ranking import batch_ranking_job
from tests.test_ranking import _fake_vector

EXPECTED_HEADER = [
    "Rank",
    "Name",
    "Email",
    "Phone",
    "Match Score",
    "Status",
    "Matched Skills",
    "Missing Skills",
    "AI Feedback",
]
TRAILING_HEADER = ["Submitted At", "Resume URL"]


class _FakeEmbedder:
    def embed(self, texts: list[str]) -> list[list[float]]:
        return [_fake_vector(t) for t in texts]


def _patch_embedder():
    return patch("app.tasks.batch_ranking.get_embedder", return_value=_FakeEmbedder())


def _create_template_with_extra_field(client: TestClient) -> dict:
    """An identity-only template would never exercise the dynamic
    template-field columns — add one extra field so the export's variable-
    width block has something real to show."""
    response = client.post(
        "/api/v1/templates",
        json={
            "template_name": f"Export Role {uuid.uuid4()}",
            "fields": [
                {
                    "field_label": "Email",
                    "field_type": "SHORT_TEXT",
                    "is_required": True,
                    "field_order": 0,
                },
                {
                    "field_label": "Full Name",
                    "field_type": "SHORT_TEXT",
                    "is_required": True,
                    "field_order": 1,
                },
                {
                    "field_label": "Years of experience",
                    "field_type": "NUMBER",
                    "is_required": True,
                    "field_order": 2,
                },
            ],
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _create_closed_job(client: TestClient) -> dict:
    template = _create_template_with_extra_field(client)
    response = client.post(
        "/api/v1/jobs",
        json={
            "job_title": "Backend Engineer <script>",
            "job_description": "Backend role needing python and sql.",
            "template_id": template["template_id"],
            "expires_at": "2027-01-01T00:00:00Z",
        },
    )
    assert response.status_code == 201, response.text
    job = response.json()
    close = client.patch(f"/api/v1/jobs/{job['job_id']}", json={"status": "CLOSED"})
    assert close.status_code == 200, close.text
    return close.json()


def _add_ranked_candidate(
    db_session: Session,
    job_id: uuid.UUID,
    email: str,
    years_field: TemplateField,
) -> Candidate:
    candidate = Candidate(
        job_id=job_id,
        full_name=email.split("@")[0],
        email=email,
        phone_number="+1-555-0100",
        submission_status="PARSED",
        resume_text=f"python developer {email}",
        submitted_at=datetime.now(UTC),
    )
    candidate.form_responses = [
        CandidateFormResponse(field_id=years_field.field_id, response_value="5")
    ]
    db_session.add(candidate)
    db_session.commit()
    db_session.refresh(candidate)
    return candidate


def _get_years_field(db_session: Session, template_id: uuid.UUID) -> TemplateField:
    return (
        db_session.query(TemplateField)
        .filter_by(template_id=template_id, field_label="Years of experience")
        .one()
    )


def _get_recruiter(db_session: Session) -> Recruiter:
    return db_session.query(Recruiter).one()


def _rank_job(authed_client: TestClient, db_session: Session, job_id: uuid.UUID) -> None:
    recruiter = _get_recruiter(db_session)
    task = BackgroundTask(
        recruiter_id=recruiter.recruiter_id, job_id=job_id, task_type=TaskType.BATCH_RANKING
    )
    db_session.add(task)
    db_session.commit()
    db_session.refresh(task)
    with _patch_embedder():
        batch_ranking_job.run(str(task.task_id))


def test_export_csv_has_bom_header_order_and_template_column(
    authed_client: TestClient, db_session: Session
) -> None:
    job = _create_closed_job(authed_client)
    job_id = uuid.UUID(job["job_id"])
    years_field = _get_years_field(db_session, uuid.UUID(job["template_id"]))
    _add_ranked_candidate(db_session, job_id, "strong@example.com", years_field)
    _rank_job(authed_client, db_session, job_id)

    response = authed_client.get(f"/api/v1/jobs/{job['job_id']}/candidates/export?format=csv")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    # script tags in job_title must never reach the header raw.
    assert "<script>" not in response.headers["content-disposition"]

    raw = response.content
    assert raw.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    # Template fields are listed in template order, including the identity
    # fields — the export repeats the raw form answers for Email/Full Name
    # alongside the canonical Name/Email columns derived from the Candidate
    # row; that's the "one column per template field" rule, no exclusion.
    assert rows[0] == EXPECTED_HEADER + ["Email", "Full Name", "Years of experience"] + (
        TRAILING_HEADER
    )
    assert rows[1][1] == "strong"  # Name
    assert rows[1][2] == "strong@example.com"  # Email
    assert rows[1][11] == "5"  # Years of experience, the dynamic column


def test_export_xlsx_default_format_matches_header_order(
    authed_client: TestClient, db_session: Session
) -> None:
    job = _create_closed_job(authed_client)
    job_id = uuid.UUID(job["job_id"])
    years_field = _get_years_field(db_session, uuid.UUID(job["template_id"]))
    _add_ranked_candidate(db_session, job_id, "a@example.com", years_field)
    _rank_job(authed_client, db_session, job_id)

    response = authed_client.get(f"/api/v1/jobs/{job['job_id']}/candidates/export")
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    header_row = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert header_row == EXPECTED_HEADER + ["Email", "Full Name", "Years of experience"] + (
        TRAILING_HEADER
    )


def test_export_is_404_cross_tenant(authed_client: TestClient, second_recruiter: Recruiter) -> None:
    job = _create_closed_job(authed_client)
    other = TestClient(authed_client.app)
    other.cookies.set(SESSION_COOKIE, create_session_cookie(str(second_recruiter.recruiter_id)))

    response = other.get(f"/api/v1/jobs/{job['job_id']}/candidates/export")
    assert response.status_code == 404


def test_export_row_cap_enforced(authed_client: TestClient, db_session: Session) -> None:
    job = _create_closed_job(authed_client)
    job_id = uuid.UUID(job["job_id"])
    years_field = _get_years_field(db_session, uuid.UUID(job["template_id"]))
    for i in range(3):
        _add_ranked_candidate(db_session, job_id, f"c{i}@example.com", years_field)
    _rank_job(authed_client, db_session, job_id)

    with patch.object(ranking_service, "EXPORT_ROW_CAP", 2):
        rows = ranking_service.export_rows(db_session, job_id)
    assert len(rows) == 2


def test_export_filename_strips_unsafe_slug_characters() -> None:
    job = Mock(apply_slug="safe-token_123")
    assert ranking_service.export_filename(job, "csv").startswith("safe-token_123-candidates-")

    job.apply_slug = "has/slash\\and space"
    filename = ranking_service.export_filename(job, "xlsx")
    assert "/" not in filename
    assert "\\" not in filename
    assert " " not in filename
    assert filename.endswith(".xlsx")
